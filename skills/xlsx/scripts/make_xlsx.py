"""Scarlet skill `xlsx` — build an .xlsx from {sheet name: rows} (openpyxl only).

Usage inside `run_code` (skill="xlsx"):

    from make_xlsx import build_xlsx
    wb = build_xlsx("budget.xlsx", {"Q1": [["Item", "Cost"], ["Rent", 1000]]})
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

#: Cap so one long cell does not make a column screen-wide.
MAX_COLUMN_WIDTH = 60


def build_xlsx(path, sheets, number_formats=None):
    """One sheet per key; row 0 is the header (bold, frozen, filter).

    A string starting with "=" is stored as a formula. `number_formats` maps
    a sheet name to {column letter: format}, e.g. {"Q1": {"B": "#,##0"}}.
    Returns the saved Workbook so callers can add charts and save again.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(list(row))
        if not rows:
            continue
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(ws.iter_cols(values_only=True), 1):
            width = max((len(str(v)) for v in column if v is not None), default=6)
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, MAX_COLUMN_WIDTH)
        for letter, fmt in (number_formats or {}).get(name, {}).items():
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    wb.save(path)
    return wb
