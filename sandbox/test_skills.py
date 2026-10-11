"""T-7.7b — script của bốn skill tự viết chạy với input mẫu ⇒ file ra mở lại được
bằng chính thư viện (docx→python-docx, xlsx→openpyxl, pptx→python-pptx, pdf→pypdf).

Chạy TRONG image `sandbox` (`make test`): đúng bộ thư viện ghim ở
`requirements.txt`, không mạng. `skills/` mount `:ro` ở `/skills` (Makefile).
"""

from __future__ import annotations

import importlib
from pathlib import Path

import pytest
from PIL import Image

SKILLS = Path("/skills")


@pytest.fixture
def load(tmp_path, monkeypatch):
    """Như `run_code`: script của skill nằm cạnh code, cwd là thư mục lượt."""
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("MPLCONFIGDIR", str(tmp_path / ".mpl"))

    def _load(skill: str, module: str):
        monkeypatch.syspath_prepend(str(SKILLS / skill / "scripts"))
        return importlib.import_module(module)

    return _load


def _png(name: str) -> str:
    Image.new("RGB", (400, 200), "white").save(name)
    return name


def test_docx_reopens_with_python_docx(load):
    from docx import Document

    make = load("docx", "make_docx")
    make.build_docx(
        "out.docx",
        [
            {"type": "heading", "text": "Mục một", "level": 1},
            {"type": "paragraph", "text": "Xin chào"},
            {"type": "bullets", "items": ["a", "b"]},
            {"type": "numbered", "items": ["c"]},
            {"type": "table", "rows": [["Tên", "Số"], ["x", 1]]},
            {"type": "image", "path": _png(".img.png"), "width_in": 3},
            {"type": "page_break"},
        ],
        title="Tiêu đề X",
    )
    document = Document("out.docx")
    assert document.paragraphs[0].text == "Tiêu đề X"
    assert document.core_properties.title == "Tiêu đề X"
    assert document.tables[0].cell(1, 1).text == "1"
    assert len(document.inline_shapes) == 1
    with pytest.raises(ValueError):
        make.build_docx("bad.docx", [{"type": "nope"}])


def test_xlsx_reopens_with_openpyxl(load):
    from openpyxl import load_workbook

    make = load("xlsx", "make_xlsx")
    make.build_xlsx(
        "out.xlsx",
        {"Chi tiêu": [["Khoản", "Tiền"], ["Nhà", 5000], ["Cộng", "=SUM(B2:B2)"]], "Trống": []},
        number_formats={"Chi tiêu": {"B": "#,##0"}},
    )
    wb = load_workbook("out.xlsx")
    ws = wb["Chi tiêu"]
    assert wb.sheetnames == ["Chi tiêu", "Trống"]
    assert ws["B3"].value == "=SUM(B2:B2)"
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws["B2"].number_format == "#,##0"


def test_pptx_reopens_with_python_pptx(load):
    from pptx import Presentation

    make = load("pptx", "make_pptx")
    make.build_pptx(
        "out.pptx",
        [
            {"title": "Kế hoạch", "subtitle": "Q4"},
            {"title": "Mục tiêu", "bullets": ["một", ["hai", 1]], "notes": "ghi chú"},
            {"title": "Biểu đồ", "image": _png(".chart.png")},
        ],
    )
    slides = list(Presentation("out.pptx").slides)
    assert [s.shapes.title.text for s in slides] == ["Kế hoạch", "Mục tiêu", "Biểu đồ"]
    body = slides[1].placeholders[1].text_frame.paragraphs
    assert [(p.text, p.level) for p in body] == [("một", 0), ("hai", 1)]
    assert slides[1].notes_slide.notes_text_frame.text == "ghi chú"
    assert any(shape.shape_type == 13 for shape in slides[2].shapes)  # PICTURE


def test_pdf_reopens_with_pypdf(load):
    import matplotlib.pyplot as plt
    from pypdf import PdfReader

    make = load("pdf", "make_pdf")
    fig = plt.figure()
    plt.plot([1, 2, 3])
    long_text = "Dòng chữ tiếng Việt có dấu, đủ dài để tự xuống hàng nhiều lần. " * 80
    make.text_pdf(".a.pdf", "Báo cáo tháng", [long_text, "Đoạn hai."], figures=[fig])
    make.merge_pdfs("out.pdf", [".a.pdf", ".a.pdf"])

    first = PdfReader(".a.pdf")
    assert first.metadata.title == "Báo cáo tháng"
    assert len(first.pages) >= 3  # chữ tràn sang trang 2 + một trang hình
    assert "Báo cáo tháng" in first.pages[0].extract_text()
    assert len(PdfReader("out.pdf").pages) == 2 * len(first.pages)
